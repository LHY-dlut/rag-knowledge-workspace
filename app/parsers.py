import io
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pdfplumber
import pymupdf as fitz
from docx import Document as WordDocument
from openpyxl import load_workbook

from app.chunking import ParsedUnit

SUPPORTED_TYPES = {".pdf", ".docx", ".xlsx", ".md", ".txt"}


def check_archive(path: Path):
    # Office files are ZIP containers: bound expanded size before parsing.
    try:
        with ZipFile(path) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Office 文件解压后过大")
    except BadZipFile as exc:
        raise ValueError("Office 文件损坏") from exc


def parse_document(path: Path, *, max_chars: int = 2_000_000, ocr: bool = False):
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise ValueError("支持 PDF / DOCX / XLSX / Markdown / TXT；旧版 DOC/XLS 请先转换")
    units: list[ParsedUnit] = []
    if suffix in {".txt", ".md"}:
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("gb18030")
        units = [ParsedUnit(content, {"location": "全文"})]
    elif suffix == ".docx":
        check_archive(path)
        doc = WordDocument(path)
        # iter_inner_content preserves paragraph/table order (python-docx 1.2).
        for index, item in enumerate(doc.iter_inner_content()):
            if hasattr(item, "rows"):
                content = "\n".join(" | ".join(c.text for c in row.cells) for row in item.rows)
            else:
                content = item.text
            if content.strip():
                units.append(ParsedUnit(content, {"location": f"段落/表格 {index + 1}"}))
    elif suffix == ".xlsx":
        check_archive(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb.worksheets:
                rows = []
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    text = " | ".join("" if cell is None else str(cell) for cell in row)
                    if text.strip(" |"):
                        rows.append(
                            ParsedUnit(
                                text,
                                {
                                    "sheet": sheet.title,
                                    "row": i,
                                    "location": f"{sheet.title} 第 {i} 行",
                                },
                            )
                        )
                units.extend(rows)
        finally:
            wb.close()
    else:
        with fitz.open(path) as pdf, pdfplumber.open(path) as tables_pdf:
            if pdf.is_encrypted:
                raise ValueError("请先解除 PDF 加密")
            for i, page in enumerate(pdf):
                text = page.get_text("text", sort=True)
                if not text.strip():
                    if not ocr:
                        raise ValueError(f"PDF 第 {i + 1} 页没有文本；启用 ENABLE_OCR 或先做 OCR")
                    try:
                        from rapidocr_onnxruntime import RapidOCR
                    except ImportError as exc:
                        raise ValueError("请安装项目的 ocr extra") from exc
                    result, _ = RapidOCR()(
                        io.BytesIO(page.get_pixmap(dpi=160).tobytes("png")).getvalue()
                    )
                    text = "\n".join(item[1] for item in (result or []))
                tables = tables_pdf.pages[i].extract_tables()
                if tables:
                    text += "\n\n[表格结构]\n" + "\n\n".join(
                        "\n".join(" | ".join(c or "" for c in row) for row in table)
                        for table in tables
                    )
                if not text.strip():
                    raise ValueError(f"PDF 第 {i + 1} 页 OCR 未提取到内容")
                units.append(ParsedUnit(text, {"page": i + 1, "location": f"第 {i + 1} 页"}))
    if sum(len(unit.content) for unit in units) > max_chars:
        raise ValueError("文档解析后超过字符上限，请拆分文件")
    if not any(x.content.strip() for x in units):
        raise ValueError("文档没有可索引文本")
    return units

import json

import httpx
import pytest
from docx import Document
from openpyxl import Workbook

from app.parsers import parse_document
from app.providers import DashScopeProvider
from app.settings import Settings


def remote_settings():
    return Settings(
        _env_file=None,
        model_provider="dashscope",
        dashscope_api_key="not-real",
        dashscope_http_base_url="https://example.test/api/v1",
        dashscope_chat_base_url="https://example.test/compatible-mode/v1",
    )


async def test_native_embedding_batch_limit_and_order():
    requests = []

    async def handle(request):
        body = json.loads(request.content)
        requests.append(body)
        texts = body["input"]["texts"]
        assert len(texts) <= 10
        assert body["parameters"] == {"dimension": 1024, "text_type": "query"}
        return httpx.Response(
            200,
            json={
                "output": {
                    "embeddings": [
                        {"text_index": i, "embedding": [float(int(t) + 1)] + [0.0] * 1023}
                        for i, t in reversed(list(enumerate(texts)))
                    ]
                }
            },
        )

    provider = DashScopeProvider(remote_settings(), httpx.MockTransport(handle))
    try:
        result = await provider.embed([str(i) for i in range(21)], "query")
        assert [v[0] for v in result] == list(range(1, 22))
        assert len(requests) == 3
    finally:
        await provider.close()


async def test_rerank_native_envelope_and_retry():
    calls = []

    async def handle(request):
        body = json.loads(request.content)
        calls.append(body)
        assert request.url.path.endswith("/services/rerank/text-rerank/text-rerank")
        assert body["input"]["query"] == "问题"
        if len(calls) == 1:
            return httpx.Response(429, json={"error": "limit"})
        return httpx.Response(
            200,
            json={
                "output": {
                    "results": [
                        {"index": 1, "relevance_score": 0.9},
                        {"index": 0, "relevance_score": 0.2},
                    ]
                }
            },
        )

    provider = DashScopeProvider(remote_settings(), httpx.MockTransport(handle))
    try:
        assert await provider.rerank("问题", ["a", "b"]) == [(1, 0.9), (0, 0.2)]
        assert len(calls) == 2
    finally:
        await provider.close()


def test_word_preserves_table_location(tmp_path):
    doc = Document()
    doc.add_paragraph("开始段落")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "报销", "7天"
    doc.add_paragraph("结束段落")
    path = tmp_path / "test.docx"
    doc.save(path)
    units = parse_document(path)
    assert [u.content for u in units] == ["开始段落", "报销 | 7天", "结束段落"]


def test_excel_and_utf8_text(tmp_path):
    wb = Workbook()
    wb.active.title = "政策"
    wb.active.append(["报销", "7天"])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    units = parse_document(path)
    assert units[0].metadata["row"] == 1 and "7天" in units[0].content
    path = tmp_path / "test.md"
    path.write_text("# 测试\n中文段落", encoding="utf-8")
    assert "中文段落" in parse_document(path)[0].content


def test_pdf_page_provenance_and_scan_failure(tmp_path):
    import pymupdf as fitz

    path = tmp_path / "test.pdf"
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Expense policy: submit within seven days.")
    pdf.save(path)
    units = parse_document(path)
    assert units[0].metadata["page"] == 1 and "seven days" in units[0].content
    blank = fitz.open()
    blank.new_page()
    path = tmp_path / "scan.pdf"
    blank.save(path)
    with pytest.raises(ValueError, match="OCR"):
        parse_document(path)
